"""
Importación del Excel EOS que envía Huawei al catálogo EOS.

Formato de cada hoja (fila 1: títulos SOFTWARE/HARDWARE, fila 2: encabezados):
    Modelo | Versión | Cantidad | SW Status | SW EOS | HW Status | HW EOS
El modelo y el EOS de hardware vienen solo en la primera fila de cada modelo
(celdas combinadas); las filas siguientes heredan esos valores.

Reglas:
  · Cada hoja de Huawei corresponde a un grupo de REDs (HOJA_A_REDS).
    La hoja Pronatel y cualquier hoja no listada se ignoran.
  · El modelo de Huawei se traduce al nombre del NCE: primero con la tabla
    cat_modelo_alias; si no hay alias, por clave_modelo() contra los modelos
    del NCE. Sin coincidencia, la fila no se importa y se reporta.
  · Cada fila se guarda en las REDs del grupo donde el NCE tiene NEs de ese
    modelo. Así el catálogo coincide con lo que reporta el NCE.
  · Los campos "Status" de Huawei no se guardan: la vigencia se calcula
    siempre desde la fecha.
  · La importación agrega y actualiza; nunca borra registros.
"""
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime

from django.db import transaction
from openpyxl import load_workbook

from inventario.models import EosHardware, EosSoftware, ModeloAlias, NceNeSnapshot, Red
from inventario.servicios import auditoria, reglas
from inventario.servicios.software import ultima_carga

DB = 'inventario'

HOJA_A_REDS = {
    'PHOTONICO': ['Fotonico', 'NFV'],
    'IPRAN': ['Acceso', 'IPRAN'],
}


class ImportacionInvalida(Exception):
    pass


@dataclass
class FilaHuawei:
    hoja: str
    fila: int
    modelo: str
    version: str
    cantidad: int
    sw_eos: date
    hw_eos: date


@dataclass
class Resultado:
    software: list = field(default_factory=list)      # cambios propuestos
    hardware: list = field(default_factory=list)
    sin_modelo: list = field(default_factory=list)    # filas cuyo modelo no existe en el NCE
    hojas_ignoradas: list = field(default_factory=list)
    cantidades: list = field(default_factory=list)    # Huawei vs NCE por hoja/modelo/versión
    errores: list = field(default_factory=list)

    def resumen(self):
        def contar(items):
            c = defaultdict(int)
            for i in items:
                c[i['accion']] += 1
            return dict(c)
        return {'software': contar(self.software), 'hardware': contar(self.hardware),
                'sin_modelo': len(self.sin_modelo), 'errores': len(self.errores)}


def _fecha(valor):
    if valor is None or valor == '':
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    try:
        return datetime.strptime(str(valor).strip(), '%d/%m/%Y').date()
    except ValueError:
        raise ImportacionInvalida(f'Fecha no reconocida: "{valor}" (se espera dd/mm/aaaa)')


def leer_excel(archivo) -> tuple[list, list]:
    """Devuelve (filas, hojas_ignoradas)."""
    wb = load_workbook(archivo, read_only=True, data_only=True)
    filas, ignoradas = [], []
    for nombre in wb.sheetnames:
        if nombre.strip().upper() not in HOJA_A_REDS:
            ignoradas.append(nombre)
            continue
        modelo = hw_eos = None
        for n, r in enumerate(wb[nombre].iter_rows(min_row=3, values_only=True), start=3):
            if not r or len(r) < 7 or not r[1]:
                continue
            if r[0]:
                modelo, hw_eos = str(r[0]).strip(), None
            if r[6]:
                hw_eos = _fecha(r[6])
            filas.append(FilaHuawei(hoja=nombre.strip().upper(), fila=n, modelo=modelo,
                                    version=str(r[1]).strip(), cantidad=int(r[2] or 0),
                                    sw_eos=_fecha(r[4]), hw_eos=hw_eos))
    if not filas:
        raise ImportacionInvalida('No se encontraron filas en las hojas Photonico / IPRAN.')
    return filas, ignoradas


def analizar(archivo) -> Resultado:
    """Compara el Excel contra el catálogo actual sin guardar nada (vista previa)."""
    filas, ignoradas = leer_excel(archivo)
    res = Resultado(hojas_ignoradas=ignoradas)

    carga = ultima_carga()
    if carga is None:
        raise ImportacionInvalida('Primero carga un NE_Report: el catálogo se valida contra los modelos del NCE.')

    alias = dict(ModeloAlias.objects.values_list('alias', 'modelo'))

    # Modelos y versiones del NCE por RED (última carga)
    nce = defaultdict(int)                       # (red, modelo, version) → NEs
    modelos_por_red = defaultdict(set)           # red → {modelo}
    por_clave = {}                               # clave_modelo → modelo NCE
    for red, modelo, version in (NceNeSnapshot.objects.all()
                                 .filter(carga=carga, ne__red__isnull=False)
                                 .values_list('ne__red__codigo', 'ne_type', 'version')):
        nce[(red, modelo, version)] += 1
        modelos_por_red[red].add(modelo)
        por_clave[reglas.clave_modelo(modelo)] = modelo

    sw_actual = {(e.red.codigo, e.modelo, e.version): e
                 for e in EosSoftware.objects.select_related('red')}
    hw_actual = {(e.red.codigo, e.modelo): e for e in EosHardware.objects.select_related('red')}
    hw_propuesto = {}

    for f in filas:
        modelo_nce = alias.get(f.modelo) or por_clave.get(reglas.clave_modelo(f.modelo))
        destino = [red for red in HOJA_A_REDS[f.hoja] if modelo_nce in modelos_por_red[red]]
        cant_nce = sum(nce[(red, modelo_nce, f.version)] for red in HOJA_A_REDS[f.hoja])
        res.cantidades.append({'hoja': f.hoja, 'modelo_huawei': f.modelo, 'modelo_nce': modelo_nce,
                               'version': f.version, 'cant_huawei': f.cantidad, 'cant_nce': cant_nce})
        if not modelo_nce or not destino:
            res.sin_modelo.append({'hoja': f.hoja, 'fila': f.fila, 'modelo': f.modelo, 'version': f.version})
            continue
        if f.sw_eos is None:
            res.errores.append({'hoja': f.hoja, 'fila': f.fila, 'error': 'Sin fecha EOS de software'})
            continue
        for red in destino:
            actual = sw_actual.get((red, modelo_nce, f.version))
            res.software.append({
                'red': red, 'modelo': modelo_nce, 'version': f.version, 'fecha_eos': f.sw_eos,
                'fecha_actual': actual.fecha_eos if actual else None,
                'accion': 'nuevo' if not actual else ('igual' if actual.fecha_eos == f.sw_eos else 'cambia'),
            })
            if f.hw_eos:
                hw_propuesto[(red, modelo_nce)] = f.hw_eos

    for (red, modelo), fecha in hw_propuesto.items():
        actual = hw_actual.get((red, modelo))
        res.hardware.append({
            'red': red, 'modelo': modelo, 'fecha_eos': fecha,
            'fecha_actual': actual.fecha_eos if actual else None,
            'accion': 'nuevo' if not actual else ('igual' if actual.fecha_eos == fecha else 'cambia'),
        })
    return res


def aplicar(archivo, usuario: str) -> Resultado:
    """Aplica la importación: crea los registros nuevos y actualiza las fechas que cambian."""
    res = analizar(archivo)
    redes = {r.codigo: r for r in Red.objects.all()}
    with transaction.atomic(using=DB):
        for c in res.software:
            if c['accion'] == 'igual':
                continue
            obj, creado = EosSoftware.objects.get_or_create(
                red=redes[c['red']], modelo=c['modelo'], version=c['version'],
                defaults={'fecha_eos': c['fecha_eos'], 'actualizado_por': usuario})
            _actualizar(obj, creado, c['fecha_eos'], usuario)
        for c in res.hardware:
            if c['accion'] == 'igual':
                continue
            obj, creado = EosHardware.objects.get_or_create(
                red=redes[c['red']], modelo=c['modelo'],
                defaults={'fecha_eos': c['fecha_eos'], 'actualizado_por': usuario})
            _actualizar(obj, creado, c['fecha_eos'], usuario)
    return res


def _actualizar(obj, creado, fecha, usuario):
    if creado:
        auditoria.registrar(usuario, 'importar', obj)
        return
    antes = auditoria.como_dict(obj)
    obj.fecha_eos, obj.actualizado_por = fecha, usuario
    obj.save()
    auditoria.registrar(usuario, 'importar', obj, antes=antes)
