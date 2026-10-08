"""
Carga del NE_Report del NCE y de la RED inicial (hoja 00_RED de la V7.5).
"""
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from inventario.models import NceCarga, NceCargaArchivo, NceNe, NceNeSnapshot, Red
from inventario.servicios import reglas
from inventario.servicios.reporte_nce import leer_reporte, parsear_fecha_nce

DB = 'inventario'
TIPO_NE_REPORT = 'NE_Report'
COLUMNAS_NE_REPORT = ['NE Name', 'NE Type', 'NE IP Address', 'NE ID', 'Software Version',
                      'Create Time', 'Running Status', 'Subnet', 'Subnet Path', 'Patch Version List']


class CargaRechazada(Exception):
    """La carga no se puede aplicar (archivo repetido, formato incorrecto, etc.)."""


def clave_ne(nombre):
    """
    Identidad de un NE: el nombre sin distinguir mayúsculas, igual que la
    collation de MySQL (la V7.5 escribe 'ASG-LIM-QUIPA' y el NCE 'ASG-LIM-Quipa').
    """
    return (nombre or '').strip().upper()


def cargar_ne_report(contenido: bytes, nombre_archivo: str) -> dict:
    """
    Registra una carga con el NE_Report: actualiza el maestro nce_ne y guarda
    la foto de cada NE dentro del alcance en nce_ne_snapshot.
    Todo en una transacción: si algo falla, no queda nada a medias.
    """
    reporte = leer_reporte(contenido, 'NE Name')
    faltan = [c for c in COLUMNAS_NE_REPORT if c not in reporte.columnas]
    if faltan:
        raise CargaRechazada(f'Al reporte le faltan columnas: {", ".join(faltan)}')

    if NceCargaArchivo.objects.filter(hash_sha256=reporte.hash_sha256).exists():
        raise CargaRechazada('Este archivo ya fue cargado antes (mismo contenido).')
    if NceCargaArchivo.objects.filter(
            carga__fecha_reporte=reporte.fecha_reporte, tipo_reporte=TIPO_NE_REPORT).exists():
        raise CargaRechazada(f'Ya existe un NE_Report para la fecha {reporte.fecha_reporte:%d/%m/%Y %H:%M}.')

    # Una fila por NE dentro del alcance (si el NCE repite un NE, vale la primera).
    filas, vistos = [], set()
    for f in reporte.filas:
        if reglas.en_alcance(f['Subnet Path']) and clave_ne(f['NE Name']) not in vistos:
            vistos.add(clave_ne(f['NE Name']))
            filas.append(f)

    with transaction.atomic(using=DB):
        carga, _ = NceCarga.objects.get_or_create(fecha_reporte=reporte.fecha_reporte)
        NceCargaArchivo.objects.create(
            carga=carga, tipo_reporte=TIPO_NE_REPORT, nombre_archivo=nombre_archivo,
            filas=len(filas), hash_sha256=reporte.hash_sha256)

        # Maestro de NEs: se crean los nuevos; los existentes conservan su RED.
        # El nombre se guarda tal como lo escribe el NCE (la fuente).
        existentes = {clave_ne(n.ne_name): n for n in NceNe.objects.all()}
        nuevos, renombrar = [], []
        for f in filas:
            ne = existentes.get(clave_ne(f['NE Name']))
            if ne is None:
                nuevos.append(NceNe(ne_name=f['NE Name'].strip(), nce_ne_id=reglas.limpiar(f['NE ID'])))
            elif ne.ne_name != f['NE Name'].strip() or ne.nce_ne_id != reglas.limpiar(f['NE ID']):
                ne.ne_name, ne.nce_ne_id = f['NE Name'].strip(), reglas.limpiar(f['NE ID'])
                ne.actualizado_en = timezone.now()
                renombrar.append(ne)
        NceNe.objects.bulk_create(nuevos, batch_size=1000)
        NceNe.objects.bulk_update(renombrar, ['ne_name', 'nce_ne_id', 'actualizado_en'], batch_size=1000)
        ids = {clave_ne(n): i for n, i in NceNe.objects.values_list('ne_name', 'id')}

        snapshots = []
        for f in filas:
            version, parche = reglas.parsear_software(f['Software Version'], f['Patch Version List'])
            snapshots.append(NceNeSnapshot(
                carga=carga, ne_id=ids[clave_ne(f['NE Name'])],
                ne_type=f['NE Type'].strip(),
                ip=reglas.limpiar(f['NE IP Address']),
                software=reglas.limpiar(f['Software Version']),
                patch_list=reglas.limpiar(f['Patch Version List']),
                version=version, parche=parche,
                estado=reglas.limpiar(f['Running Status']),
                creado_en_nce=parsear_fecha_nce(f['Create Time']),
                subnet_path=f['Subnet Path'].strip(),
                subnet=reglas.limpiar(f['Subnet']),
            ))
        NceNeSnapshot.objects.bulk_create(snapshots, batch_size=1000)

        carga.estado = 'ok'
        carga.fin = timezone.now()
        carga.save(update_fields=['estado', 'fin'])

    return {
        'carga_id': carga.id,
        'fecha_reporte': reporte.fecha_reporte,
        'filas_reporte': len(reporte.filas),
        'nes_en_alcance': len(filas),
        'nes_nuevos': len(nuevos),
        'sin_red': NceNe.objects.filter(snapshots__carga=carga, red__isnull=True).count(),
    }


def cargar_red_desde_excel(ruta_excel, hoja=None, usuario='carga_inicial') -> dict:
    """
    Carga inicial de la RED de cada NE desde la hoja 00_RED de la V7.5
    (columnas: NE Name, RED). Los NEs que no existen se crean en el maestro.
    """
    wb = load_workbook(ruta_excel, read_only=True, data_only=True)
    nombre_hoja = hoja or next((n for n in wb.sheetnames if n.upper().startswith('00_RED')), None)
    if not nombre_hoja:
        raise CargaRechazada(f'No hay una hoja que empiece con "00_RED". Hojas: {", ".join(wb.sheetnames)}')

    redes = {r.codigo.upper(): r for r in Red.objects.all()}
    filas = list(wb[nombre_hoja].iter_rows(min_row=2, values_only=True))
    desconocidas, asignaciones = set(), {}
    for fila in filas:
        if not fila or not fila[0] or not fila[1]:
            continue
        red = redes.get(str(fila[1]).strip().upper())
        if red is None:
            desconocidas.add(str(fila[1]).strip())
            continue
        asignaciones[clave_ne(str(fila[0]))] = (str(fila[0]).strip(), red)
    if desconocidas:
        raise CargaRechazada(f'RED no reconocida en la hoja {nombre_hoja}: {", ".join(sorted(desconocidas))}')

    with transaction.atomic(using=DB):
        existentes = {clave_ne(n.ne_name): n for n in NceNe.objects.all()}
        crear, actualizar = [], []
        for clave, (ne_name, red) in asignaciones.items():
            ne = existentes.get(clave)
            if ne is None:
                crear.append(NceNe(ne_name=ne_name, red=red, red_confirmada=True, actualizado_por=usuario))
            elif ne.red_id != red.id or not ne.red_confirmada:
                ne.red, ne.red_confirmada, ne.actualizado_por = red, True, usuario
                ne.actualizado_en = timezone.now()
                actualizar.append(ne)
        NceNe.objects.bulk_create(crear, batch_size=1000)
        NceNe.objects.bulk_update(actualizar, ['red', 'red_confirmada', 'actualizado_por', 'actualizado_en'], batch_size=1000)

    return {'hoja': nombre_hoja, 'filas': len(asignaciones), 'creados': len(crear), 'actualizados': len(actualizar)}
