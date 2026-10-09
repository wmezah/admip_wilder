"""
Importación de catálogos desde Excel con vista previa (Part Numbers, Seriales,
RED de NEs). Misma lógica para la carga inicial desde la V7.5 y para el botón
"Importar Excel" de la página:

  · analizar(): lee el Excel y lo compara con la base, sin guardar nada.
  · aplicar():  crea los nuevos y actualiza los que cambian. Nunca borra.

Cada catálogo acepta los encabezados de la V7.5 y los del Excel que exporta
la página (ver COLUMNAS). Si una clave se repite en el archivo, vale la
primera fila (igual que BUSCARV en Excel) y las demás se reportan.
"""
from dataclasses import dataclass, field

from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from inventario.models import NceNe, PartNumber, Red, SerialPartNumber
from inventario.servicios import auditoria
from inventario.servicios.carga_ne import clave_ne

DB = 'inventario'


class ImportacionInvalida(Exception):
    pass


def _texto(v):
    return '' if v is None else str(v).strip()


def _booleano(v):
    """Enable de la V7.5: True/False o el texto TRUE/FALSE. Vacío = inventariable."""
    if v is None or _texto(v) == '':
        return True
    if isinstance(v, bool):
        return v
    t = _texto(v).upper()
    if t in ('TRUE', 'SI', 'SÍ', '1', 'VERDADERO'):
        return True
    if t in ('FALSE', 'NO', '0', 'FALSO'):
        return False
    raise ValueError(f'valor "{v}" no es Sí/No')


ELEMENTOS_PN = {e for e, _ in PartNumber.ELEMENTOS}


def _elemento(v):
    t = _texto(v)
    mapa = {e.upper(): e for e in ELEMENTOS_PN if e}
    if t == '':
        return ''
    if t.upper() not in mapa:
        raise ValueError(f'elemento "{t}" no válido ({", ".join(sorted(mapa.values()))})')
    return mapa[t.upper()]


# Para cada catálogo: hoja preferida, encabezados aceptados por campo y conversión.
COLUMNAS = {
    'part_numbers': {
        'hoja': 'BOMCODE',
        'campos': {
            'pn': (['PN(BOM Code) Item', 'PN'], _texto),
            'descripcion': (['Description Item', 'Descripción', 'Descripcion'], _texto),
            'inventariable': (['Enable', 'Inventariable'], _booleano),
            'observacion': (['Remarks', 'Observación', 'Observacion'], _texto),
            'elemento': (['Element', 'Elemento'], _elemento),
        },
        'clave': 'pn',
    },
    'seriales': {
        'hoja': 'SERIAL',
        'campos': {
            'serial': (['Serial', 'SN'], _texto),
            'pn': (['PN(BOM Code) Power', 'PN'], _texto),
            'comentario': (['Comentario'], _texto),
        },
        'clave': 'serial',
    },
    'red_nes': {
        'hoja': '00_RED',
        'campos': {
            'ne_name': (['NE Name', 'NE'], _texto),
            'red': (['RED'], _texto),
        },
        'clave': 'ne_name',
    },
}


@dataclass
class Vista:
    nuevos: list = field(default_factory=list)
    cambian: list = field(default_factory=list)
    iguales: int = 0
    errores: list = field(default_factory=list)
    repetidos: list = field(default_factory=list)
    hoja: str = ''

    def resumen(self):
        return {'nuevos': len(self.nuevos), 'cambian': len(self.cambian), 'iguales': self.iguales,
                'errores': len(self.errores), 'repetidos': len(self.repetidos), 'hoja': self.hoja}


def _leer(archivo, catalogo):
    cfg = COLUMNAS[catalogo]
    wb = load_workbook(archivo, read_only=True, data_only=True)
    hoja = next((n for n in wb.sheetnames if n.upper().startswith(cfg['hoja'].upper())), wb.sheetnames[0])
    filas = wb[hoja].iter_rows(values_only=True)
    encabezado = [_texto(c) for c in next(filas, [])]
    indices = {}
    for campo, (nombres, _) in cfg['campos'].items():
        idx = next((encabezado.index(n) for n in nombres if n in encabezado), None)
        if idx is not None:
            indices[campo] = idx
    if cfg['clave'] not in indices:
        raise ImportacionInvalida(
            f'La hoja "{hoja}" no tiene la columna {cfg["campos"][cfg["clave"]][0][0]}. '
            f'Encabezados encontrados: {", ".join(e for e in encabezado if e)}')
    registros, vistos, errores, repetidos = [], set(), [], []
    for n, fila in enumerate(filas, start=2):
        if not fila or all(v is None for v in fila):
            continue
        reg = {}
        try:
            for campo, idx in indices.items():
                reg[campo] = cfg['campos'][campo][1](fila[idx] if idx < len(fila) else None)
        except ValueError as e:
            errores.append({'fila': n, 'error': str(e)})
            continue
        clave = reg[cfg['clave']]
        if not clave:
            continue
        k = clave.upper()
        if k in vistos:
            repetidos.append({'fila': n, 'clave': clave})
            continue
        vistos.add(k)
        registros.append((n, reg))
    return hoja, registros, errores, repetidos


def _actuales(catalogo):
    if catalogo == 'part_numbers':
        return {p.pn.upper(): p for p in PartNumber.objects.all()}
    if catalogo == 'seriales':
        return {s.serial.upper(): s for s in SerialPartNumber.objects.all()}
    return {clave_ne(n.ne_name): n for n in NceNe.objects.select_related('red')}


def _comparable(catalogo, obj):
    if catalogo == 'red_nes':
        return {'red': obj.red.codigo if obj.red else '', 'confirmada': obj.red_confirmada}
    campos = [c for c in COLUMNAS[catalogo]['campos'] if c != COLUMNAS[catalogo]['clave']]
    return {c: getattr(obj, c) for c in campos}


def analizar(archivo, catalogo):
    hoja, registros, errores, repetidos = _leer(archivo, catalogo)
    actuales = _actuales(catalogo)
    redes = {r.codigo.upper(): r.codigo for r in Red.objects.all()}
    vista = Vista(errores=errores, repetidos=repetidos, hoja=hoja)
    for n, reg in registros:
        if catalogo == 'red_nes':
            if not reg.get('red'):
                continue            # fila sin RED: no hay nada que importar
            red = redes.get(reg.get('red', '').upper())
            if red is None:
                vista.errores.append({'fila': n, 'error': f'RED "{reg.get("red")}" no existe'})
                continue
            reg = {'ne_name': reg['ne_name'], 'red': red, 'confirmada': True}
        clave = reg[COLUMNAS[catalogo]['clave']]
        obj = actuales.get(clave_ne(clave) if catalogo == 'red_nes' else clave.upper())
        if obj is None:
            vista.nuevos.append(reg)
            continue
        antes = _comparable(catalogo, obj)
        nuevos_valores = {k: v for k, v in reg.items() if k in antes}
        if any(antes[k] != v for k, v in nuevos_valores.items()):
            vista.cambian.append({**reg, 'antes': antes})
        else:
            vista.iguales += 1
    return vista


def aplicar(archivo, catalogo, usuario):
    vista = analizar(archivo, catalogo)
    actuales = _actuales(catalogo)
    redes = {r.codigo: r for r in Red.objects.all()}
    with transaction.atomic(using=DB):
        for reg in vista.nuevos + vista.cambian:
            reg = {k: v for k, v in reg.items() if k != 'antes'}
            if catalogo == 'part_numbers':
                obj = actuales.get(reg['pn'].upper()) or PartNumber(pn=reg['pn'])
                antes = auditoria.como_dict(obj) if obj.pk else None
                for k, v in reg.items():
                    setattr(obj, k, v)
            elif catalogo == 'seriales':
                obj = actuales.get(reg['serial'].upper()) or SerialPartNumber(serial=reg['serial'])
                antes = auditoria.como_dict(obj) if obj.pk else None
                for k, v in reg.items():
                    setattr(obj, k, v)
            else:
                obj = actuales.get(clave_ne(reg['ne_name'])) or NceNe(ne_name=reg['ne_name'])
                antes = auditoria.como_dict(obj) if obj.pk else None
                obj.red, obj.red_confirmada = redes[reg['red']], True
                obj.actualizado_en = timezone.now()
            obj.actualizado_por = usuario
            obj.save()
            auditoria.registrar(usuario, 'importar', obj, antes=antes)
    return vista


# ─── Exportar (mismos encabezados que acepta la importación) ─────────────────

ENCABEZADOS_EXPORTAR = {
    'part_numbers': ['PN', 'Descripción', 'Elemento', 'Inventariable', 'Observación'],
    'seriales': ['Serial', 'PN', 'Comentario'],
    'red_nes': ['NE Name', 'RED', 'Confirmada'],
}


def exportar(catalogo):
    from io import BytesIO
    from openpyxl import Workbook
    wb = Workbook(write_only=True)
    ws = wb.create_sheet({'part_numbers': 'BOMCODE', 'seriales': 'SERIAL', 'red_nes': '00_RED'}[catalogo])
    ws.append(ENCABEZADOS_EXPORTAR[catalogo])
    if catalogo == 'part_numbers':
        for p in PartNumber.objects.all():
            ws.append([p.pn, p.descripcion, p.elemento, 'Sí' if p.inventariable else 'No', p.observacion])
    elif catalogo == 'seriales':
        for s in SerialPartNumber.objects.all():
            ws.append([s.serial, s.pn, s.comentario])
    else:
        for n in NceNe.objects.select_related('red').order_by('ne_name'):
            ws.append([n.ne_name, n.red.codigo if n.red else '', 'Sí' if n.red_confirmada else 'No'])
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()
