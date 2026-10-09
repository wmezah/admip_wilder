"""
Inventario de una carga: componentes del NCE → ítems (reglas V7.5) → cambios
contra la carga anterior. También la carga base desde la hoja INVENTARIO de
la V7.5, para tener contra qué comparar la primera carga.
"""
from collections import defaultdict
from datetime import datetime

from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from inventario.models import (
    InvCambio, InvItem, NceCarga, NceComponente, NceNe, NceNeSnapshot, PartNumber, Red, SerialPartNumber,
)
from inventario.servicios import reglas
from inventario.servicios.carga_ne import CargaRechazada, clave_ne
from inventario.servicios.reglas_inventario import Catalogos, ItemPN, calcular_inventario

DB = 'inventario'


# ─── Componentes ─────────────────────────────────────────────────────────────

def pn_nce(valor):
    """PN tal como lo usa la V7.5: '--' se conserva (tiene regla propia); los demás vacíos → ''."""
    v = (valor or '').strip()
    return v if v == '--' else reglas.limpiar(v)


def guardar_componentes(carga, reportes, nes_en_alcance):
    """
    Guarda en nce_componente las filas de Subrack, Board, Subcard y OpticalModule
    de los NEs en alcance. `nes_en_alcance` es un set de clave_ne.
    """
    def del_alcance(ne):
        return clave_ne(ne) in nes_en_alcance

    filas = []
    for f in reportes['Subrack_Report'].filas:
        if del_alcance(f['NE']):
            filas.append(NceComponente(
                carga=carga, tipo='frame', ne_name=f['NE'].strip(),
                pn=pn_nce(f['PN(BOM Code/Item)']), sn=reglas.limpiar(f['SN(Bar Code)']),
                descripcion=reglas.limpiar(f['Description'])))
    for f in reportes['Board_Report'].filas:
        if del_alcance(f['NE']):
            filas.append(NceComponente(
                carga=carga, tipo='board', ne_name=f['NE'].strip(), nombre=f['Board Name'].strip(),
                slot=f['Slot ID'].strip(), pn=pn_nce(f['PN(BOM Code/Item)']),
                sn=reglas.limpiar(f['SN(Bar Code)']), descripcion=reglas.limpiar(f['Description'])))
    for f in reportes['Subcard_Report'].filas:
        if del_alcance(f['NE']):
            filas.append(NceComponente(
                carga=carga, tipo='subboard', ne_name=f['NE'].strip(), nombre=f['Subboard Type'].strip(),
                slot=f['Slot Number'].strip(), subslot=f['Subslot Number'].strip(),
                pn=pn_nce(f['PN(BOM Code/Item)']), sn=reglas.limpiar(f['SN(Bar Code)']),
                descripcion=reglas.limpiar(f['Subboard Description'])))
    for f in reportes['OpticalModule_Information'].filas:
        if del_alcance(f['NE Name']):
            # Los transceivers se guardan con sus valores originales: sus reglas
            # distinguen "-", "--" y vacío (ver reglas_inventario.pn_transceiver).
            filas.append(NceComponente(
                carga=carga, tipo='transceiver', ne_name=f['NE Name'].strip(),
                nombre=f['Optical/Electrical Type'].strip(), puerto=f['Port Name'].strip(),
                pn=f['PN(BOM Code/Item)'].strip(), sn=f['Serial No.'].strip(),
                vendor_pn=f['Vendor PN'].strip(), port_custom=f['Port Custom Column'].strip()))
    NceComponente.objects.bulk_create(filas, batch_size=2000)
    return len(filas)


# ─── Catálogos ───────────────────────────────────────────────────────────────

def catalogos_actuales():
    return Catalogos(
        part_numbers={p.pn: ItemPN(descripcion=p.descripcion, inventariable=p.inventariable)
                      for p in PartNumber.objects.all()},
        seriales=dict(SerialPartNumber.objects.values_list('serial', 'pn')),
    )


# ─── Inventario ──────────────────────────────────────────────────────────────

def _entradas(carga):
    """Arma las entradas de reglas_inventario.calcular_inventario desde la base."""
    nes = {}
    for s in NceNeSnapshot.objects.filter(carga=carga, ne__red__isnull=False).select_related('ne__red'):
        if reglas.es_planta_fisica(s.ne_type, s.subnet_path):
            nes[clave_ne(s.ne.ne_name)] = {'ne_type': s.ne_type, 'red': s.ne.red.codigo, 'ne_id': s.ne_id,
                                           'ne_name': s.ne.ne_name}

    frames, boards, subboards, transceivers = {}, [], [], []
    for c in NceComponente.objects.filter(carga=carga).order_by('id'):
        ne = clave_ne(c.ne_name)
        if c.tipo == 'frame':
            # Una fila por NE: se prefiere la que tiene SN (V7.5).
            if ne not in frames or (not frames[ne]['sn'] and c.sn):
                frames[ne] = {'pn': c.pn, 'sn': c.sn, 'descripcion': c.descripcion}
        elif c.tipo == 'board':
            boards.append({'ne': ne, 'nombre': c.nombre, 'slot': c.slot, 'pn': c.pn, 'sn': c.sn,
                           'descripcion': c.descripcion})
        elif c.tipo == 'subboard':
            subboards.append({'ne': ne, 'nombre': c.nombre, 'slot': c.slot, 'subslot': c.subslot,
                              'pn': c.pn, 'sn': c.sn, 'descripcion': c.descripcion})
        else:
            transceivers.append({'ne': ne, 'tipo': c.nombre, 'puerto': c.puerto, 'pn': c.pn,
                                 'vendor_pn': c.vendor_pn, 'port_custom': c.port_custom, 'serial': c.sn})
    # La V7.5 recorre los chasis en orden de NE (hoja 02_Frame ordenada)
    frames = dict(sorted(frames.items()))
    return nes, frames, boards, subboards, transceivers


def calcular_items(carga):
    """Calcula y guarda los ítems del inventario de la carga. Devuelve el conteo por elemento."""
    nes, frames, boards, subboards, transceivers = _entradas(carga)
    items = calcular_inventario(nes, frames, boards, subboards, transceivers, catalogos_actuales())
    redes = dict(Red.objects.values_list('codigo', 'id'))
    InvItem.objects.bulk_create([
        InvItem(carga=carga, ne_id=nes[i['ne']]['ne_id'], red_id=redes[i['red']], modelo=i['modelo'],
                pn_chasis=i['pn_chasis'], elemento=i['elemento'], nombre=i['nombre'][:150],
                sr=i['sr'], b=i['b'], s=i['s'], p=i['p'][:20], pn=i['pn'][:60], sn=i['sn'][:80],
                descripcion=i['descripcion'][:500])
        for i in items], batch_size=2000)
    conteo = defaultdict(int)
    for i in items:
        conteo[i['elemento']] += 1
    return dict(conteo)


# ─── Cambios ─────────────────────────────────────────────────────────────────

def _clave_item(i):
    """Con SN: elemento + SN. Sin SN (algunas subtarjetas): elemento + NE + posición."""
    sn = (i.sn or '').strip().upper()
    if sn:
        return (i.elemento, sn)
    return (i.elemento, i.ne_id, i.nombre, i.b, i.s, i.p)


def _posicion(i):
    """
    Texto para mostrar: chasis → su nombre · tarjeta → 'POWER 5' (el nombre ya
    trae el slot) · subtarjeta → 'ETH_10xGF_CARD 10/1' · transceiver → 'SFPPLUS 0/2/9'.
    """
    if i.elemento in ('Chasis', 'Board'):
        return i.nombre
    partes = [x for x in (i.b, i.s, i.p) if x and x != '.']
    return f'{i.nombre} {"/".join(partes)}'.strip()


def _lugar(i):
    """Lo que define un movimiento: NE y posición física (no el nombre, que el NCE a veces reetiqueta)."""
    return (i.ne_id, i.b, i.s, i.p)


def carga_anterior(carga):
    return (NceCarga.objects.filter(estado='ok', fecha_reporte__lt=carga.fecha_reporte, items__isnull=False)
            .order_by('-fecha_reporte').distinct().first())


def detectar_cambios(carga):
    """
    Compara los ítems de la carga con los de la carga anterior:
      alta: está ahora y no antes · baja: estaba antes y no ahora ·
      movimiento: mismo ítem (por SN) en otro NE u otra posición física.
    """
    anterior = carga_anterior(carga)
    if anterior is None:
        return None
    campos = ('id', 'ne_id', 'ne__ne_name', 'red__codigo', 'elemento', 'nombre', 'b', 's', 'p', 'pn', 'sn', 'descripcion')

    def agrupar(c):
        grupos = defaultdict(list)
        for i in InvItem.objects.filter(carga=c).select_related('ne', 'red').only(*campos):
            grupos[_clave_item(i)].append(i)
        return grupos

    antes, ahora = agrupar(anterior), agrupar(carga)
    cambios = []

    def nuevo(tipo, i, a=None):
        base = i or a
        cambios.append(InvCambio(
            carga=carga, fecha=carga.fecha_reporte, fecha_anterior=anterior.fecha_reporte, tipo=tipo,
            elemento=base.elemento, red=base.red.codigo, pn=base.pn, sn=base.sn, descripcion=base.descripcion,
            ne_antes=a.ne.ne_name if a else '', pos_antes=_posicion(a) if a else '',
            ne_despues=i.ne.ne_name if i else '', pos_despues=_posicion(i) if i else ''))

    for clave in set(antes) | set(ahora):
        a_lista, i_lista = antes.get(clave, []), ahora.get(clave, [])
        for a, i in zip(a_lista, i_lista):
            if _lugar(a) != _lugar(i):
                nuevo('movimiento', i, a)
        for i in i_lista[len(a_lista):]:
            nuevo('alta', i)
        for a in a_lista[len(i_lista):]:
            nuevo('baja', None, a)
    InvCambio.objects.bulk_create(cambios, batch_size=2000)
    resumen = defaultdict(int)
    for c in cambios:
        resumen[c.tipo] += 1
    return dict(resumen)


def recalcular(carga):
    """Vuelve a calcular el inventario y los cambios de una carga (tras corregir un catálogo)."""
    with transaction.atomic(using=DB):
        InvItem.objects.filter(carga=carga).delete()
        InvCambio.objects.filter(carga=carga).delete()
        conteo = calcular_items(carga)
        cambios = detectar_cambios(carga)
    return conteo, cambios


# ─── Carga base desde la V7.5 ────────────────────────────────────────────────

def cargar_base_v75(ruta_excel, fecha):
    """
    Registra la hoja INVENTARIO de la V7.5 como una carga 'base_v75' en la
    fecha indicada. Sirve solo como punto de comparación para la primera carga.
    """
    if NceCarga.objects.filter(origen='base_v75').exists():
        raise CargaRechazada('Ya existe una carga base de la V7.5.')
    wb = load_workbook(ruta_excel, read_only=True, data_only=True)
    if 'INVENTARIO' not in wb.sheetnames:
        raise CargaRechazada('El Excel no tiene la hoja INVENTARIO.')
    redes = {r.codigo.upper(): r.id for r in Red.objects.all()}
    filas = [r for r in wb['INVENTARIO'].iter_rows(min_row=2, values_only=True)
             if r and str(r[0]).strip().upper() == 'TRUE' and r[1] and str(r[1]).upper() in redes]
    # (la columna E de la V7.5 tiene TRUE como valor lógico y también como texto)

    with transaction.atomic(using=DB):
        fecha_dt = timezone.make_aware(datetime.combine(fecha, datetime.min.time()))
        carga = NceCarga.objects.create(fecha_reporte=fecha_dt, origen='base_v75', estado='procesando',
                                        mensaje='Hoja INVENTARIO del Excel V7.5')
        existentes = {clave_ne(n): i for n, i in NceNe.objects.values_list('ne_name', 'id')}
        faltan = {clave_ne(str(r[2])): (str(r[2]).strip(), redes[str(r[1]).upper()])
                  for r in filas if clave_ne(str(r[2])) not in existentes}
        NceNe.objects.bulk_create([NceNe(ne_name=n, red_id=red, red_confirmada=True, actualizado_por='base V7.5')
                                   for n, red in faltan.values()], batch_size=1000)
        ids = {clave_ne(n): i for n, i in NceNe.objects.values_list('ne_name', 'id')}
        txt = lambda v: '' if v is None else str(v).strip()
        InvItem.objects.bulk_create([
            InvItem(carga=carga, ne_id=ids[clave_ne(str(r[2]))], red_id=redes[str(r[1]).upper()],
                    pn_chasis=txt(r[3]), modelo=txt(r[4]), elemento=txt(r[6]), nombre=txt(r[7])[:150],
                    sr=txt(r[8]), b=txt(r[9]), s=txt(r[10]), p=txt(r[11])[:20],
                    pn=txt(r[12])[:60], sn=txt(r[13])[:80], descripcion=txt(r[14])[:500])
            for r in filas], batch_size=2000)
        carga.estado, carga.fin = 'ok', timezone.now()
        carga.save(update_fields=['estado', 'fin'])
    return {'carga_id': carga.id, 'items': len(filas), 'nes_creados': len(faltan)}
