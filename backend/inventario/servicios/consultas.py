"""
Consultas de la página Planta instalada Huawei: resumen, equipos integrados
por mes, pendientes y el Excel completo. Todo sale de la última carga.
"""
import io
from collections import Counter, defaultdict

from django.db.models import Count
from django.utils import timezone
from openpyxl import Workbook

from inventario.models import InvCambio, InvItem, NceComponente, NceNeSnapshot, Red
from inventario.servicios import reglas
from inventario.servicios.inventario import carga_anterior, catalogos_actuales
from inventario.servicios.reglas_inventario import pn_transceiver, sn_transceiver
from inventario.servicios.software import ultima_carga

ELEMENTOS = ['Chasis', 'Board', 'SubBoard', 'Transceiver']


def _fecha(dt):
    return timezone.localtime(dt) if dt else None


def _snapshots(carga, red=None):
    """Fotos de NEs de planta física con RED (los que forman el inventario)."""
    qs = NceNeSnapshot.objects.filter(carga=carga, ne__red__isnull=False).select_related('ne__red')
    if red:
        qs = qs.filter(ne__red__codigo=red)
    return [s for s in qs if reglas.es_planta_fisica(s.ne_type, s.subnet_path)]


def _tipo_transceiver(descripcion):
    """'SFP+, 10Gbps-1310nm-LC-10Km(0.009mm)' (primeras dos partes de la descripción)."""
    return ','.join((descripcion or 'Sin descripción').split(',')[:2]).strip()


def resumen(red=None):
    carga = ultima_carga()
    if carga is None:
        return None
    items = InvItem.objects.filter(carga=carga)
    if red:
        items = items.filter(red__codigo=red)
    snaps = _snapshots(carga, red)

    por_elemento = dict(items.values_list('elemento').annotate(n=Count('id')))
    matriz_raw = (InvItem.objects.filter(carga=carga).values('elemento', 'red__codigo').annotate(n=Count('id')))
    redes = list(Red.objects.values_list('codigo', flat=True))
    matriz = {e: {r: 0 for r in redes} for e in ELEMENTOS}
    for m in matriz_raw:
        matriz[m['elemento']][m['red__codigo']] = m['n']

    modelos = defaultdict(Counter)
    for modelo, r in items.filter(elemento='Chasis').values_list('modelo', 'red__codigo'):
        modelos[modelo][r] += 1
    top_modelos = sorted(({'modelo': m, 'total': sum(c.values()), 'red': c.most_common(1)[0][0]}
                          for m, c in modelos.items()), key=lambda x: -x['total'])[:10]

    tipos = Counter(_tipo_transceiver(d) for d in
                    items.filter(elemento='Transceiver').values_list('descripcion', flat=True))
    software = Counter((reglas.clave_version(s.version) and s.version[:11]) or 'Sin versión' for s in snaps)
    top_sw = software.most_common(5)
    otras = sum(software.values()) - sum(n for _, n in top_sw)

    cambios = InvCambio.objects.filter(carga=carga)
    if red:
        cambios = cambios.filter(red=red)
    cambios_tabla = {e: {'alta': 0, 'baja': 0, 'movimiento': 0} for e in ELEMENTOS}
    for c in cambios.values('elemento', 'tipo').annotate(n=Count('id')):
        cambios_tabla[c['elemento']][c['tipo']] = c['n']
    anterior = carga_anterior(carga)

    return {
        'carga': {'id': carga.id, 'fecha_reporte': _fecha(carga.fecha_reporte), 'origen': carga.origen},
        'carga_anterior': ({'id': anterior.id, 'fecha_reporte': _fecha(anterior.fecha_reporte),
                            'origen': anterior.origen} if anterior else None),
        'red': red or 'all',
        'kpis': {
            'nes': len(snaps),
            'no_normal': sum(1 for s in snaps if s.estado != 'Normal'),
            **{e.lower(): por_elemento.get(e, 0) for e in ELEMENTOS},
            'items': sum(por_elemento.values()),
        },
        'matriz': [{'elemento': e, **matriz[e]} for e in ELEMENTOS],
        'top_modelos': top_modelos,
        'top_transceivers': [{'tipo': t, 'total': n} for t, n in tipos.most_common(8)],
        'software': [{'version': v, 'total': n} for v, n in top_sw] + ([{'version': 'Otras', 'total': otras}] if otras else []),
        'cambios': {
            'altas': sum(v['alta'] for v in cambios_tabla.values()),
            'bajas': sum(v['baja'] for v in cambios_tabla.values()),
            'movimientos': sum(v['movimiento'] for v in cambios_tabla.values()),
            'por_elemento': [{'elemento': e, **cambios_tabla[e]} for e in ELEMENTOS],
        },
        'pendientes': conteo_pendientes(carga),
    }


# ─── Pendientes ──────────────────────────────────────────────────────────────

def pendientes(carga=None):
    """
    · NEs sin RED y NEs con RED sugerida sin confirmar.
    · PN sin descripción: ni el NCE ni el catálogo de Part Numbers la tienen.
    · Transceivers con SN pero sin PN (no entran al inventario hasta registrar
      su SN en el catálogo de seriales).
    """
    carga = carga or ultima_carga()
    if carga is None:
        return None
    snaps = NceNeSnapshot.objects.filter(carga=carga).select_related('ne__red')
    sin_red, por_confirmar, del_inventario = [], [], set()
    for s in snaps:
        if not reglas.es_planta_fisica(s.ne_type, s.subnet_path):
            continue
        if s.ne.red is not None:
            del_inventario.add(s.ne.ne_name.strip().upper())
        fila = {'ne_id': s.ne_id, 'ne': s.ne.ne_name, 'modelo': s.ne_type, 'subnet_path': s.subnet_path,
                'red': s.ne.red.codigo if s.ne.red else ''}
        if s.ne.red is None:
            sin_red.append(fila)
        elif not s.ne.red_confirmada:
            por_confirmar.append(fila)

    sin_descripcion = (InvItem.objects.filter(carga=carga, descripcion='')
                       .values('pn', 'elemento').annotate(cantidad=Count('id')).order_by('-cantidad'))

    cat = catalogos_actuales()
    sin_pn = []
    for c in NceComponente.objects.filter(carga=carga, tipo='transceiver'):
        if c.ne_name.strip().upper() not in del_inventario:
            continue
        sn = sn_transceiver(c.sn)
        t = {'pn': c.pn, 'vendor_pn': c.vendor_pn, 'port_custom': c.port_custom}
        if sn and not pn_transceiver(t, sn, cat):
            sin_pn.append({'ne': c.ne_name, 'puerto': c.puerto, 'tipo': c.nombre, 'sn': sn})

    return {'nes_sin_red': sin_red, 'red_por_confirmar': por_confirmar,
            'pn_sin_descripcion': list(sin_descripcion), 'transceivers_sin_pn': sin_pn}


def conteo_pendientes(carga):
    p = pendientes(carga)
    return {k: len(v) for k, v in p.items()}


# ─── Equipos integrados por mes ──────────────────────────────────────────────

def integrados(anio, red=None, mes=None):
    """
    NEs por mes de alta en el NCE ("Create Time") dentro de un año.
    Solo cuenta NEs que siguen en la red en la última carga.
    """
    carga = ultima_carga()
    if carga is None:
        return None
    snaps = NceNeSnapshot.objects.filter(carga=carga, creado_en_nce__isnull=False).select_related('ne__red')
    filas = [s for s in snaps if reglas.es_planta_fisica(s.ne_type, s.subnet_path)
             and (not red or (s.ne.red and s.ne.red.codigo == red))]
    anios = Counter(timezone.localtime(s.creado_en_nce).year for s in filas)
    del_anio = [s for s in filas if timezone.localtime(s.creado_en_nce).year == anio]
    redes = list(Red.objects.values_list('codigo', flat=True)) + ['Sin RED']
    meses = [{'mes': m, **{r: 0 for r in redes}, 'total': 0} for m in range(1, 13)]
    modelos = Counter()
    for s in del_anio:
        m = timezone.localtime(s.creado_en_nce).month
        r = s.ne.red.codigo if s.ne.red else 'Sin RED'
        meses[m - 1][r] += 1
        meses[m - 1]['total'] += 1
        modelos[s.ne_type] += 1
    detalle = None
    if mes:
        detalle = sorted(({'ne': s.ne.ne_name, 'modelo': s.ne_type, 'red': s.ne.red.codigo if s.ne.red else '',
                           'integrado': timezone.localtime(s.creado_en_nce).date(), 'subnet': s.subnet,
                           'ip': s.ip}
                          for s in del_anio if timezone.localtime(s.creado_en_nce).month == mes),
                         key=lambda x: x['integrado'])
    return {
        'anio': anio, 'red': red or 'all', 'total': len(del_anio),
        'anios_disponibles': sorted(anios), 'por_anio': dict(sorted(anios.items())),
        'meses': meses, 'top_modelos': [{'modelo': m, 'total': n} for m, n in modelos.most_common(5)],
        'detalle_mes': detalle,
    }


# ─── Excel completo ──────────────────────────────────────────────────────────

ENCABEZADO_INVENTARIO = ['RED', 'Group', 'BOM Code Group', 'Description Group', 'Element', 'Name',
                         'SR', 'B', 'S', 'P', 'PN(BOM Code) Item', 'SN(Bar Code) Item', 'Description Item']


def excel_completo(red=None):
    """Excel con el formato de la hoja INVENTARIO de la V7.5, más NEs, cambios y pendientes."""
    carga = ultima_carga()
    if carga is None:
        return None, None
    wb = Workbook(write_only=True)

    ws = wb.create_sheet('INVENTARIO')
    ws.append(ENCABEZADO_INVENTARIO)
    items = InvItem.objects.filter(carga=carga).order_by('red__orden', 'ne__ne_name', 'id')
    if red:
        items = items.filter(red__codigo=red)
    for fila in items.values_list('red__codigo', 'ne__ne_name', 'pn_chasis', 'modelo', 'elemento', 'nombre',
                                  'sr', 'b', 's', 'p', 'pn', 'sn', 'descripcion').iterator(chunk_size=5000):
        ws.append(list(fila))

    ws = wb.create_sheet('NE')
    ws.append(['NE Name', 'RED', 'RED confirmada', 'NE Type', 'IP', 'Software', 'Versión', 'Parche',
               'Estado', 'Alta en NCE', 'Subnet Path'])
    for s in _snapshots(carga, red):
        ws.append([s.ne.ne_name, s.ne.red.codigo, 'Sí' if s.ne.red_confirmada else 'Sugerida', s.ne_type, s.ip,
                   s.software, s.version, s.parche, s.estado,
                   timezone.localtime(s.creado_en_nce).replace(tzinfo=None) if s.creado_en_nce else None,
                   s.subnet_path])

    ws = wb.create_sheet('Cambios')
    ws.append(['Tipo', 'Element', 'RED', 'PN', 'SN', 'Descripción', 'NE antes', 'Posición antes',
               'NE ahora', 'Posición ahora'])
    cambios = InvCambio.objects.filter(carga=carga)
    if red:
        cambios = cambios.filter(red=red)
    for c in cambios:
        ws.append([c.get_tipo_display(), c.elemento, c.red, c.pn, c.sn, c.descripcion,
                   c.ne_antes, c.pos_antes, c.ne_despues, c.pos_despues])

    p = pendientes(carga)
    ws = wb.create_sheet('Pendientes')
    ws.append(['Tipo', 'NE / PN', 'Modelo / Elemento', 'Detalle'])
    for f in p['nes_sin_red']:
        ws.append(['NE sin RED', f['ne'], f['modelo'], f['subnet_path']])
    for f in p['red_por_confirmar']:
        ws.append(['RED sugerida por confirmar', f['ne'], f['modelo'], f'{f["red"]} · {f["subnet_path"]}'])
    for f in p['pn_sin_descripcion']:
        ws.append(['PN sin descripción', f['pn'], f['elemento'], f'{f["cantidad"]} ítems'])
    for f in p['transceivers_sin_pn']:
        ws.append(['Transceiver sin PN', f['ne'], f['tipo'], f'{f["puerto"]} · SN {f["sn"]}'])

    salida = io.BytesIO()
    wb.save(salida)
    fecha = timezone.localtime(carga.fecha_reporte)
    nombre = f'Inventario_Huawei_{red or "Todas"}_{fecha:%Y-%m-%d}.xlsx'
    return salida.getvalue(), nombre
